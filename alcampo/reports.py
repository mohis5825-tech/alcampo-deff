import os
import re
import smtplib
import ssl
import tempfile
from email.message import EmailMessage
from pathlib import Path

from .core import ConfigError


def export_excel(rows, path: Path, metadata):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Resultados"
    sheet.append(["Cuenta", "Estado", "Saldo EUR", "Comprobado UTC", "Detalle"])
    ordered = sorted(rows, key=lambda row: (row["saldo_centimos"] is None, -(row["saldo_centimos"] or 0), row["cuenta"]))
    for row in ordered:
        sheet.append([row["cuenta"], row["estado"],
                      row["saldo_centimos"] / 100 if row["saldo_centimos"] is not None else None,
                      row["comprobado"], row["detalle"]])
        for col in (1, 2, 4, 5):
            sheet.cell(sheet.max_row, col).data_type = "s"
        sheet.cell(sheet.max_row, 3).number_format = '#,##0.00 "€"'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in (("A", 38), ("B", 30), ("C", 18), ("D", 28), ("E", 65)):
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    summary = book.create_sheet("Resumen")
    for key, value in metadata.items():
        summary.append([key, value])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 65
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(fd)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        book.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


def export_invalid_excel(rows, path: Path, metadata):
    """Guarda aparte las cuentas sin un saldo válido/publicable."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Cuentas excluidas"
    sheet.append(["Cuenta", "Estado", "Comprobado UTC", "Detalle"])
    for row in sorted(rows, key=lambda item: (item["estado"], item["cuenta"])):
        sheet.append([row["cuenta"], row["estado"], row["comprobado"], row["detalle"]])
        for col in range(1, 5):
            sheet.cell(sheet.max_row, col).data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in (("A", 38), ("B", 30), ("C", 28), ("D", 80)):
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    summary = book.create_sheet("Resumen")
    for key, value in metadata.items():
        summary.append([key, value])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 65
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(fd)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        book.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


def telegram_config(config):
    token = os.environ.get("BOT_TOKEN", "").strip()
    chat = str(os.environ.get("CHAT_ID") or config.get("telegram_chat_id", "")).strip()
    if not re.fullmatch(r"\d+:[A-Za-z0-9_-]{20,}", token) or not re.fullmatch(r"-?\d+", chat):
        raise ConfigError("Telegram necesita BOT_TOKEN y un CHAT_ID numérico. Consulta CONFIGURACION.md")
    return token, chat


def send_excel(path, config, caption):
    import requests

    token, chat = telegram_config(config)
    try:
        with Path(path).open("rb") as document:
            response = requests.post(
                f"https://api.telegram.org/bot{token}/sendDocument",
                data={"chat_id": chat, "caption": caption[:1024]},
                files={"document": (Path(path).name, document,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                timeout=(10, 60), allow_redirects=False,
            )
        response.raise_for_status()
        if response.status_code != 200 or response.json().get("ok") is not True:
            raise ValueError
    except Exception:
        raise ConfigError("Telegram no confirmó el envío. El Excel está guardado; revisa token, chat y acceso del bot. No se reenvió automáticamente") from None


def email_config():
    """Lee SMTP desde el entorno sin imprimir credenciales."""
    recipient = os.environ.get("EMAIL_TO", "").strip()
    username = os.environ.get("EMAIL_USERNAME", "").strip()
    password = os.environ.get("EMAIL_PASSWORD", "")
    host = os.environ.get("SMTP_HOST", "").strip() or "smtp.gmail.com"
    security = os.environ.get("SMTP_SECURITY", "").strip().lower() or "ssl"
    raw_port = os.environ.get("SMTP_PORT", "").strip() or "465"
    try:
        port = int(raw_port)
    except ValueError:
        port = 0
    valid_address = r"[^\s@]+@[^\s@]+\.[^\s@]+"
    if not re.fullmatch(valid_address, recipient) or not re.fullmatch(valid_address, username):
        raise ConfigError("El correo necesita EMAIL_TO y EMAIL_USERNAME válidos")
    if not password or not host or port not in range(1, 65536) or security not in {"ssl", "starttls"}:
        raise ConfigError("El correo necesita SMTP_HOST, SMTP_PORT, SMTP_SECURITY y EMAIL_PASSWORD válidos")
    return recipient, username, password, host, port, security


def send_email(path, caption):
    """Envía el Excel como adjunto mediante SMTP."""
    recipient, username, password, host, port, security = email_config()
    message = EmailMessage()
    message["From"] = username
    message["To"] = recipient
    message["Subject"] = "Informe Alcampo"
    message.set_content(caption)
    with Path(path).open("rb") as document:
        message.add_attachment(
            document.read(),
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=Path(path).name,
        )
    try:
        context = ssl.create_default_context()
        if security == "ssl":
            with smtplib.SMTP_SSL(host, port, context=context, timeout=60) as server:
                server.login(username, password)
                server.send_message(message)
        else:
            with smtplib.SMTP(host, port, timeout=60) as server:
                server.ehlo()
                server.starttls(context=context)
                server.ehlo()
                server.login(username, password)
                server.send_message(message)
    except Exception:
        raise ConfigError("El correo no confirmó el envío; revisa EMAIL_TO, SMTP y la contraseña de aplicación. El Excel está guardado") from None
